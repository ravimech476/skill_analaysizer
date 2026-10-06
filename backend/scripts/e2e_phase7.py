"""End-to-end test for Phase 7 (Excel uploads for marks / staff / skills, exports and the dashboard).

Runs against the demo data (seed_demo.py) and puts back every mark and skill level it changes, so it can be
re-run. Staff it creates get random employee codes and are deactivated at the end. Usage:
  API_URL=http://localhost:8090/api/v1 python scripts/e2e_phase7.py
"""
import io
import os
import random
import sys

import openpyxl
import requests

B = os.environ.get("API_URL", "http://localhost:8080/api/v1")
ENV = os.path.join(os.path.dirname(__file__), "..", ".env")
ADMIN_PW = next(l.split("=", 1)[1].strip() for l in open(ENV) if l.startswith("ADMIN_PASSWORD="))
SFX = str(random.randint(1000, 9999))
PASS = FAIL = 0
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def check(name, resp, status):
    global PASS, FAIL
    ok = resp.status_code == status
    PASS += ok
    FAIL += not ok
    print(f"  {'ok  ' if ok else 'FAIL'} {name} ({resp.status_code}){'' if ok else ' ' + resp.text[:300]}")
    return resp.json().get("data") if resp.headers.get("content-type", "").startswith("application/json") else None


def expect(name, cond, info=""):
    global PASS, FAIL
    PASS += bool(cond)
    FAIL += not cond
    print(f"  {'ok  ' if cond else 'FAIL'} {name} {info}")


def login(u, p):
    r = requests.post(f"{B}/auth/login", json={"username": u, "password": p})
    r.raise_for_status()
    return {"Authorization": "Bearer " + r.json()["data"]["access_token"]}


def get(h, path, **kw):
    return requests.get(B + path, headers=h, **kw)


def data(h, path):
    r = get(h, path)
    if r.status_code != 200:
        sys.exit(f"GET {path}: {r.status_code} {r.text}")
    return r.json()["data"]


def book(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def xlsx_bytes(rows, info=None, sheet="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    if info:
        wi = wb.create_sheet("Info")
        for k, v in info.items():
            wi.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def post_file(h, path, content, form=None, name="upload.xlsx"):
    return requests.post(B + path, headers=h, files={"file": (name, content, XLSX)}, data=form or {})


A = login("admin", ADMIN_PW)
PRIYA = login("priya", "Staff@1234")    # staff + placement officer, CSE, allocated CS3301
MEENA = login("meena", "Staff@1234")    # staff, CSE, incharge CSE II-A
KAVYA = login("kavya", "Staff@1234")    # CSE HOD
SURESH = login("suresh", "Staff@1234")  # ECE HOD
ARUN = login("25cs001", "Student@123")
PARENT = login("p9200000001", "Parent@123")

cls = next(c for c in data(A, "/classes") if c["label"] == "CSE II-A")
subjects = {s["code"]: s["id"] for s in data(A, "/subjects?all=true")}
exams = {e["code"]: e["id"] for e in data(A, "/exam-types?all=true")}
students = {s["register_no"]: s for s in data(A, "/students?page_size=200&status=all")}
KEY = {"class_id": cls["id"], "semester_id": cls["current_semester_id"], "subject_id": subjects["CS3301"], "exam_type_id": exams["IA1"], "attempt_no": 1}
Q = "&".join(f"{k}={v}" for k, v in KEY.items())

# ---------------------------------------------------------------- marks upload
print("marks template")
r = get(PRIYA, f"/marks/entry/template?{Q}")
check("allocated staff downloads the marks template", r, 200)
expect("template is an xlsx attachment", r.headers.get("content-type") == XLSX and "attachment" in r.headers.get("content-disposition", ""))
wb = book(r)
expect("template has Marks + Info sheets", wb.sheetnames[:2] == ["Marks", "Info"], wb.sheetnames)
rows = list(wb["Marks"].iter_rows(values_only=True))
expect("header row", [str(x).lower() for x in rows[0][:3]] == ["register_no*", "name", "marks*"], rows[0])
orig = {r[0]: r[2] for r in rows[1:]}
expect("roster pre-filled with existing marks", orig.get("25CS001") is not None and float(orig["25CS001"]) == 46, orig)
info = {r[0]: r[1] for r in wb["Info"].iter_rows(values_only=True)}
expect("Info pins the exam", str(info.get("subject_id")) == str(subjects["CS3301"]) and str(info.get("exam_type_id")) == str(exams["IA1"]), info)
check("template needs the key", get(PRIYA, "/marks/entry/template?class_id=1"), 400)

original_file = r.content  # uploading the untouched template restores the marks later

print("marks upload: guards")
form = {k: str(v) for k, v in KEY.items()}
check("other department's HOD cannot upload", post_file(SURESH, "/marks/entry/upload", original_file, form), 403)
check("student cannot upload", post_file(ARUN, "/marks/entry/upload", original_file, form), 403)
check("upload needs the key fields", post_file(PRIYA, "/marks/entry/upload", original_file, {}), 400)
check("non-xlsx rejected", requests.post(B + "/marks/entry/upload", headers=PRIYA, files={"file": ("m.csv", b"a,b", "text/csv")}, data=form), 400)
check("missing marks column rejected", post_file(PRIYA, "/marks/entry/upload", xlsx_bytes([["register_no", "name"], ["25CS001", "x"]]), form), 400)
ia2 = dict(form, exam_type_id=str(exams["IA2"]))
r = post_file(PRIYA, "/marks/entry/upload", original_file, ia2)
check("IA1 template uploaded as IA2 is refused", r, 400)
expect("mismatch message names the problem", "different class, subject or exam" in r.text, r.text[:120])

print("marks upload: row errors (nothing saved)")
bad = xlsx_bytes([["register_no", "name", "marks"], ["25CS001", "Arun", 47], ["25CS002", "Anjali", "abc"],
                  ["25CS003", "Bharath", 51], ["25CS999", "Nobody", 10], ["25CS001", "Arun again", 40], ["25CS004", "Divya", ""]])
res = check("upload with bad rows", post_file(PRIYA, "/marks/entry/upload", bad, form), 200)
job = res["job"]
expect("not saved", res["saved"] is False)
expect("4 row errors", job["failed_rows"] == 4, [e["message"] for e in job["errors"]])
expect("success 0 (all-or-nothing)", job["success_rows"] == 0)
expect("error rows numbered from the sheet", sorted(e["row"] for e in job["errors"]) == [3, 4, 5, 6], [e["row"] for e in job["errors"]])
expect("job type is marks", job["upload_type"] == "marks")
entry = data(PRIYA, f"/marks/entry?{Q}")
m = {row["register_no"]: row for row in entry["rows"]}
expect("valid row not saved either", m["25CS001"]["marks_obtained"] == 46, m["25CS001"]["marks_obtained"])

print("marks upload: dry run")
good = xlsx_bytes([["register_no", "name", "marks"], ["25CS001", "Arun", 47.5], ["25cs003", "Bharath", "ab"], ["25CS004", "Divya", ""]],
                  info={k: v for k, v in KEY.items()}, sheet="Marks")
res = check("dry run", post_file(PRIYA, "/marks/entry/upload", good, dict(form, dry_run="true")), 200)
expect("dry run validates 2 rows, saves nothing", res["saved"] is False and res["job"]["success_rows"] == 2 and res["job"]["dry_run"], res["job"])
expect("marks unchanged after dry run", data(PRIYA, f"/marks/entry?{Q}")["rows"][0]["marks_obtained"] == 46)

print("marks upload: save")
res = check("upload saves", post_file(PRIYA, "/marks/entry/upload", good, form), 200)
expect("saved", res["saved"] is True and res["job"]["success_rows"] == 2)
m = {row["register_no"]: row for row in res["sheet"]["rows"]}
expect("decimal mark stored", m["25CS001"]["marks_obtained"] == 47.5, m["25CS001"])
expect("AB → absent (case-insensitive register no)", m["25CS003"]["is_absent"] is True, m["25CS003"])
expect("blank cell left unchanged", m["25CS004"]["marks_obtained"] is not None and m["25CS004"]["marks_obtained"] == float(orig["25CS004"]), m["25CS004"])
notes = data(ARUN, "/notifications?page_size=5")
notes = notes if isinstance(notes, list) else notes.get("items", notes)
expect("student notified of the change", any("marks" in n["title"] and "CS3301" in n["title"] for n in notes), [n["title"] for n in notes][:3])
res = check("uploading the original template restores the marks", post_file(PRIYA, "/marks/entry/upload", original_file, form), 200)
m = {row["register_no"]: row for row in res["sheet"]["rows"]}
expect("restored", m["25CS001"]["marks_obtained"] == 46 and not m["25CS003"]["is_absent"], (m["25CS001"]["marks_obtained"], m["25CS003"]["is_absent"]))
r = get(A, "/bulk-upload/jobs?upload_type=marks&page_size=50")
jobs = r.json()["data"]
expect("marks jobs listed by type", r.status_code == 200 and jobs and all(j["upload_type"] == "marks" for j in jobs), len(jobs))
r = get(PRIYA, f"/bulk-upload/jobs/{res['job']['id']}")
check("uploader can open the job", r, 200)
r = get(PRIYA, f"/bulk-upload/jobs/{job['id']}/errors.xlsx")
check("error report download", r, 200)
expect("error report lists the 4 rows", book(r)["Errors"].max_row >= 6, book(r)["Errors"].max_row)

# ---------------------------------------------------------------- class result sheet
print("class result sheet xlsx")
sq = f"class_id={cls['id']}&semester_id={cls['current_semester_id']}&exam_type_id={exams['SEM']}"
r = get(MEENA, f"/marks/sheet.xlsx?{sq}")
check("incharge downloads the result sheet", r, 200)
wb = book(r)
expect("Results + Subject summary sheets", wb.sheetnames == ["Results", "Subject summary"], wb.sheetnames)
ws = wb["Results"]
expect("title names the class", "CSE II-A" in str(ws["A1"].value), ws["A1"].value)
hdr = [c.value for c in ws[4]]
expect("subject columns in header", {"CS3301", "CS3311", "MA3301"} <= set(hdr), hdr)
body = {row[0]: row for row in ws.iter_rows(min_row=5, values_only=True)}
expect("Arun's end-sem mark with grade", "94" in str(body["25CS001"][hdr.index("CS3301")]), body["25CS001"])
expect("Divya fails one subject", str(body["25CS004"][hdr.index("Result")]).startswith("Fail"), body["25CS004"])
summ = list(wb["Subject summary"].iter_rows(min_row=5, values_only=True))
expect("summary has a row per subject", len(summ) == 3, len(summ))
check("JSON sheet still works", get(MEENA, f"/marks/sheet?{sq}"), 200)
check("student cannot download the result sheet", get(ARUN, f"/marks/sheet.xlsx?{sq}"), 403)
check("other department's HOD cannot", get(SURESH, f"/marks/sheet.xlsx?{sq}"), 403)
check("result sheet needs the key", get(MEENA, "/marks/sheet.xlsx"), 400)

# ---------------------------------------------------------------- mark statement pdf
print("mark statement pdf")
arun, anjali, lakshmi = students["25CS001"]["id"], students["25CS002"]["id"], students["25EC001"]["id"]
r = get(ARUN, f"/marks/students/{arun}/statement.pdf")
check("student downloads own statement", r, 200)
expect("is a PDF", r.content[:5] == b"%PDF-" and r.headers.get("content-type") == "application/pdf", r.content[:8])
expect("file named by register no", "25CS001" in r.headers.get("content-disposition", ""))
check("student cannot download a classmate's", get(ARUN, f"/marks/students/{anjali}/statement.pdf"), 404)
check("parent downloads a child's statement", get(PARENT, f"/marks/students/{anjali}/statement.pdf"), 200)
check("parent cannot download another student's", get(PARENT, f"/marks/students/{students['25CS003']['id']}/statement.pdf"), 404)
check("staff downloads own department's", get(MEENA, f"/marks/students/{arun}/statement.pdf"), 200)
check("staff cannot download another department's", get(MEENA, f"/marks/students/{lakshmi}/statement.pdf"), 404)
r = get(A, f"/marks/students/{lakshmi}/statement.pdf")
check("statement with no marks still renders", r, 200)
expect("empty statement is a PDF", r.content[:5] == b"%PDF-")

# ---------------------------------------------------------------- students export
print("students xlsx")
r = get(A, "/students/export.xlsx?lifecycle=studying")
check("admin exports students", r, 200)
ws = book(r)["Students"]
regs = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
expect("all studying students exported", len(regs) >= 12 and "25EC001" in regs, len(regs))
hdr = [c.value for c in ws[1]]
row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "25CS001")
expect("parent columns filled", row[hdr.index("Parent")] == "Kumar Ramasamy" and row[hdr.index("Parent mobile")] == "9200000001", row)
expect("register numbers kept as text", isinstance(row[0], str))
r = get(MEENA, "/students/export.xlsx")
regs = [row[0] for row in book(r)["Students"].iter_rows(min_row=2, values_only=True)]
expect("staff export limited to own department", regs and all(x.startswith("2") and "CS" in x for x in regs), regs)
r = get(A, f"/students/export.xlsx?class_id={cls['id']}")
expect("class filter applies", len(list(book(r)["Students"].iter_rows(min_row=2))) == 4)
check("student cannot export", get(ARUN, "/students/export.xlsx"), 403)
check("parent cannot export", get(PARENT, "/students/export.xlsx"), 403)

# ---------------------------------------------------------------- placement exports
print("placement xlsx")
r = get(A, "/placements.xlsx")
check("admin exports placements", r, 200)
wb = book(r)
rows = list(wb["Placements"].iter_rows(min_row=3, values_only=True))
expect("Arun's Zoho offer exported", any(r_[0] == "25CS001" and r_[4] == "Zoho" for r_ in rows), rows[:3])
expect("company summary sheet", "By company" in wb.sheetnames)
r = get(SURESH, "/placements.xlsx")
rows = [x for x in book(r)["Placements"].iter_rows(min_row=3, values_only=True) if x[0]]
expect("ECE HOD sees no CSE placements", all(not str(x[0]).startswith("25CS") for x in rows), rows)
check("student cannot export placements", get(ARUN, "/placements.xlsx"), 403)

roles = data(A, "/job-roles?page_size=50")
roles = roles if isinstance(roles, list) else roles.get("items", [])
role = next(x for x in roles if x["company_name"] == "TCS")
check("analyze TCS", requests.post(f"{B}/job-roles/{role['id']}/analyze", headers=A), 200)
r = get(A, f"/job-roles/{role['id']}/matches.xlsx")
check("ranking export", r, 200)
wb = book(r)
ws = wb["Ranking"]
expect("ranking title names the drive", "TCS" in str(ws["A1"].value), ws["A1"].value)
hdr = [c.value for c in ws[4]]
rk = list(ws.iter_rows(min_row=5, values_only=True))
expect("ranked rows present", len(rk) > 0 and hdr[0] == "Rank", (len(rk), hdr[:3]))
elig = [x for x in rk if x[hdr.index("Eligible")] == "Yes"]
expect("eligible rows ranked 1..n", [x[0] for x in elig] == list(range(1, len(elig) + 1)), [x[0] for x in elig])
expect("ineligible rows carry reasons", all(x[hdr.index("Why not eligible")] for x in rk if x[hdr.index("Eligible")] == "No"))
expect("required skills sheet", wb["Required skills"].max_row >= 4)
r = get(A, f"/job-roles/{role['id']}/matches.xlsx?eligible_only=true")
expect("eligible_only filter", all(x[hdr.index("Eligible")] == "Yes" for x in book(r)["Ranking"].iter_rows(min_row=5, values_only=True)))
check("ranking of unknown role", get(A, "/job-roles/999999/matches.xlsx"), 404)
check("student cannot export a ranking", get(ARUN, f"/job-roles/{role['id']}/matches.xlsx"), 403)
check("JSON matches still work", get(A, f"/job-roles/{role['id']}/matches"), 200)
check("JSON matches of unknown role still 404", get(A, "/job-roles/999999/matches"), 404)

# ---------------------------------------------------------------- staff upload
print("staff upload")
r = get(A, "/bulk-upload/staff/template")
check("staff template", r, 200)
wb = book(r)
expect("template sheets", wb.sheetnames == ["Staff", "Instructions", "Departments"], wb.sheetnames)
check("plain staff cannot download the staff template", get(MEENA, "/bulk-upload/staff/template"), 403)
hdr = ["employee_code*", "name*", "department_code", "designation", "qualification", "roles", "mobile", "email", "gender", "dob", "joined_on"]
c1, c2 = f"T7A{SFX}", f"T7B{SFX}"
rows = [hdr,
        [c1, "Test Staff One", "CSE", "Assistant Professor", "M.E.", "staff", None, f"t7a{SFX}@x.edu", "female", "1990-01-15", "2021-06-01"],
        [c2, "Test Staff Two", "ece", "Professor", "Ph.D.", "staff, hod", None, None, "male", None, "01-07-2019"],
        [f"T7C{SFX}", "Bad Dept", "XYZ", None, None, None, None, None, None, None, None],
        [f"T7D{SFX}", "Bad Role", "CSE", None, None, "principal", None, None, None, None, None],
        ["EMP001", "Duplicate Priya", "CSE", None, None, None, None, None, None, None, None],
        [f"T7E{SFX}", "Bad mobile", "CSE", None, None, None, "12345", None, None, None, None],
        [f"T7F{SFX}", "Future dob", "CSE", None, None, None, None, None, None, "2090-01-01", None],
        [None, None, None, None, None, None, None, None, None, None, None]]
content = xlsx_bytes(rows)
res = check("staff dry run", post_file(A, "/bulk-upload/staff", content, {"dry_run": "true", "default_password": "Welcome@123"}), 201)
expect("dry run: 2 ok, 5 failed", res["success_rows"] == 2 and res["failed_rows"] == 5, [e["message"] for e in res["errors"]])
msgs = " | ".join(e["message"] for e in res["errors"])
expect("errors explain themselves", "department_code" in msgs and "roles" in msgs and "employee code" in msgs.lower() and "Mobile" in msgs and "dob" in msgs, msgs)
found = data(A, f"/staff?search=T7A{SFX}")
found = found if isinstance(found, list) else found.get("items", [])
expect("dry run created nothing", len(found) == 0)
check("short default password refused", post_file(A, "/bulk-upload/staff", content, {"default_password": "short"}), 400)
check("plain staff cannot upload staff", post_file(MEENA, "/bulk-upload/staff", content), 403)
res = check("staff upload", post_file(A, "/bulk-upload/staff", content, {"default_password": "Welcome@123"}), 201)
expect("2 created", res["success_rows"] == 2 and res["upload_type"] == "staff")
lst = requests.get(f"{B}/staff?search=T7", headers=A).json()["data"]
new = {s["employee_code"]: s for s in lst if s["employee_code"] in (c1, c2)}
expect("both staff exist", len(new) == 2, list(new))
expect("department + roles applied", new[c2]["department_name"].startswith("Electronics") and set(new[c2]["roles"]) == {"staff", "hod"}, new.get(c2))
expect("dates parsed (DD-MM-YYYY)", new[c2]["joined_on"] == "2019-07-01", new[c2]["joined_on"])
r = requests.post(f"{B}/auth/login", json={"username": c1.lower(), "password": "Welcome@123"})
check("new staff logs in with the default password", r, 200)
res = check("re-upload reports duplicates", post_file(A, "/bulk-upload/staff", xlsx_bytes(rows[:3])), 201)
expect("both rows now duplicates", res["success_rows"] == 0 and res["failed_rows"] == 2)
for s in new.values():
    requests.patch(f"{B}/staff/{s['id']}/status", headers=A, json={"is_active": False})

# ---------------------------------------------------------------- skills upload
print("skills upload")
r = get(MEENA, f"/bulk-upload/skills/template?class_id={cls['id']}")
check("skills template pre-filled for a class", r, 200)
wb = book(r)
rws = list(wb["Skills"].iter_rows(min_row=2, values_only=True))
expect("one row per student of the class", len(rws) == 4 and rws[0][0] == "25CS001", [x[0] for x in rws])
expect("skill list sheet", "Skill list" in wb.sheetnames and wb["Skill list"].max_row > 5)
check("student cannot download the skills template", get(ARUN, "/bulk-upload/skills/template"), 403)

before = {s["name"]: s["proficiency"] for s in data(A, f"/students/{arun}/skills")}
skill_names = [row[0] for row in wb["Skill list"].iter_rows(min_row=2, values_only=True)]
sk1 = "Python" if "Python" in skill_names else skill_names[0]
sk2 = next(n for n in skill_names if n not in before and n != sk1)
lvl1 = 5 if before.get(sk1) != 5 else 4
rows = [["register_no", "name", "skill", "level", "source", "remarks"],
        ["25CS001", "Arun", sk1.lower(), lvl1, "certification", "e2e"],
        ["25CS001", "Arun", sk2, 2, None, None],
        ["25CS002", "Anjali", "Not A Skill", 3, None, None],
        ["25CS003", "Bharath", sk1, 7, None, None],
        ["25EC001", "Lakshmi", sk1, 3, None, None],
        ["25CS001", "Arun dup", sk1, 3, None, None],
        ["25CS004", "Divya", sk1, None, None, None],
        ["25CS004", "Divya", sk1, 3, "hackathon", None],
        ["99XX999", "Nobody", sk1, 3, None, None]]
content = xlsx_bytes(rows)
res = check("skills dry run (staff)", post_file(MEENA, "/bulk-upload/skills", content, {"dry_run": "true"}), 201)
expect("dry run: 2 ok, 6 failed (blank level skipped)", res["success_rows"] == 2 and res["failed_rows"] == 6, [e["message"] for e in res["errors"]])
msgs = " | ".join(e["message"] for e in res["errors"])
expect("skill errors explain themselves", all(k in msgs for k in ["Unknown skill", "1 to 5", "another department", "Duplicate", "source", "No student"]), msgs)
expect("dry run changed nothing", {s["name"]: s["proficiency"] for s in data(A, f"/students/{arun}/skills")} == before)
check("student cannot upload skills", post_file(ARUN, "/bulk-upload/skills", content), 403)
res = check("skills upload", post_file(MEENA, "/bulk-upload/skills", content), 201)
after = {s["name"]: s for s in data(A, f"/students/{arun}/skills")}
expect("level updated (skill name case-insensitive)", after[sk1]["proficiency"] == lvl1 and after[sk1]["source"] == "certification", after.get(sk1))
expect("new skill added with default source", after[sk2]["proficiency"] == 2 and after[sk2]["source"] == "assessment", after.get(sk2))
notes = data(ARUN, "/notifications?page_size=5")
notes = notes if isinstance(notes, list) else notes.get("items", notes)
expect("one combined skill notice", notes and notes[0]["title"] == "Skills updated" and sk2 in notes[0]["body"] and sk1 in notes[0]["body"], notes[0] if notes else None)
res = check("admin may upload across departments", post_file(A, "/bulk-upload/skills", xlsx_bytes([rows[0], ["25EC001", "Lakshmi", sk1, 3, None, None]]), {"dry_run": "true"}), 201)
expect("admin: other department accepted", res["success_rows"] == 1)
# restore Arun's skills
skill_ids = {row["name"]: row["id"] for row in data(A, "/skills?all=true")}
if sk1 in before:
    requests.put(f"{B}/students/{arun}/skills", headers=A, json={"skill_id": skill_ids[sk1], "proficiency": before[sk1]})
else:
    requests.delete(f"{B}/students/{arun}/skills/{skill_ids[sk1]}", headers=A)
requests.delete(f"{B}/students/{arun}/skills/{skill_ids[sk2]}", headers=A)
expect("skills restored", {s["name"]: s["proficiency"] for s in data(A, f"/students/{arun}/skills")} == before)

# ---------------------------------------------------------------- dashboard
print("dashboard")
d = check("admin dashboard", get(A, "/reports/dashboard"), 200)
expect("whole-college scope", d["scope_department"] is None)
h = d["headline"]
expect("headline numbers", h["students"] >= 12 and h["staff"] >= 4 and h["placed"] >= 1 and h["avg_cgpa"], h)
expect("exam list and default exam", d["exams"] and d["exam"] is not None, d["exam"])
d2 = check("dashboard for the end-sem exam", get(A, f"/reports/dashboard?exam_type_id={exams['SEM']}"), 200)
cp = next((x for x in d2["pass_by_class"] if x["class_label"] == "CSE II-A"), None)
expect("CSE II-A pass % (11 of 12 end-sem papers passed)", cp and cp["appeared"] == 12 and cp["passed"] == 11 and cp["pass_percent"] == 91.67, cp)
expect("all-clear students (3 of 4)", cp and cp["all_clear"] == 3 and cp["all_clear_percent"] == 75, cp)
expect("CGPA buckets", len(d["cgpa_distribution"]) == 6 and sum(b["total"] for b in d["cgpa_distribution"]) + d["cgpa_not_graded"] == h["students"],
       [b["total"] for b in d["cgpa_distribution"]])
expect("skill gaps sorted by coverage", d["skill_gaps"] and all(d["skill_gaps"][i]["coverage_percent"] <= d["skill_gaps"][i + 1]["coverage_percent"] for i in range(len(d["skill_gaps"]) - 1)),
       [g["coverage_percent"] for g in d["skill_gaps"]])
expect("placement by batch", any(b["batch"] == "2025-2029" and b["placed"] >= 1 for b in d["placement_by_batch"]), d["placement_by_batch"])
expect("12 months of placement trend", len(d["placement_by_month"]) == 12)
dm = check("staff dashboard", get(MEENA, "/reports/dashboard"), 200)
expect("staff scoped to own department", dm["scope_department"] == "CSE" and dm["departments"] == ["CSE"], (dm["scope_department"], dm["departments"]))
ds = check("ECE HOD asking for CSE still gets ECE", get(SURESH, f"/reports/dashboard?department_id={students['25CS001']['department_id']}"), 200)
expect("HOD scope forced", ds["scope_department"] == "ECE" and all(x["department_code"] == "ECE" for x in ds["pass_by_class"]), ds["scope_department"])
dp = check("placement officer can filter a department", get(PRIYA, f"/reports/dashboard?department_id={students['25EC001']['department_id']}"), 200)
expect("filter applied", dp["scope_department"] == "ECE" and dp["headline"]["students"] == 3, dp["headline"]["students"])
check("unknown department", get(A, "/reports/dashboard?department_id=999999"), 400)
check("student has no dashboard", get(ARUN, "/reports/dashboard"), 403)
check("parent has no dashboard", get(PARENT, "/reports/dashboard"), 403)
me = requests.get(f"{B}/auth/me", headers=MEENA).json()["data"]
expect("staff hold report.view", "report.view" in me.get("permissions", []), me.get("permissions", [])[:5])

print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
