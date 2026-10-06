"""End-to-end test for Phase 2 (masters, classes, curriculum, staff, students, parents, bulk upload).

Needs the API on :8080 (OTP_DEBUG=true) and the Phase 1 demo users (admin, priya). Creates data with a
random suffix so it can be re-run. Usage: python scripts/e2e_phase2.py
"""
import io
import os
import random
import sys

import openpyxl
import requests

B = os.environ.get("API_URL", "http://localhost:8080/api/v1")
ENV = os.path.join(os.path.dirname(__file__), "..", ".env")
ADMIN_PW = next(l.split("=", 1)[1].strip() for l in open(ENV) if l.startswith("ADMIN_PASSWORD="))
SFX = str(random.randint(100, 999))
PASS = FAIL = 0


def check(name, resp, status):
    global PASS, FAIL
    ok = resp.status_code == status
    PASS += ok
    FAIL += not ok
    detail = "" if ok else f" expected {status} got {resp.status_code}: {resp.text[:300]}"
    print(f"  {'ok  ' if ok else 'FAIL'} {name} ({resp.status_code}){detail}")
    return resp.json() if resp.headers.get("content-type", "").startswith("application/json") else None


def expect(name, cond, info=""):
    global PASS, FAIL
    PASS += bool(cond)
    FAIL += not cond
    print(f"  {'ok  ' if cond else 'FAIL'} {name} {info}")


def login(username, password):
    r = requests.post(f"{B}/auth/login", json={"username": username, "password": password})
    return {"Authorization": "Bearer " + r.json()["data"]["access_token"]}


def otp_login(identifier):
    sent = requests.post(f"{B}/auth/otp/request", json={"identifier": identifier}).json()["data"]
    r = requests.post(f"{B}/auth/otp/verify", json={"identifier": identifier, "otp": sent["debug_otp"]})
    return {"Authorization": "Bearer " + r.json()["data"]["access_token"]}


A = login("admin", ADMIN_PW)
get = lambda p, h=A, **kw: requests.get(B + p, headers=h, **kw)
post = lambda p, body, h=A: requests.post(B + p, json=body, headers=h)
put = lambda p, body, h=A: requests.put(B + p, json=body, headers=h)
patch = lambda p, body, h=A: requests.patch(B + p, json=body, headers=h)
delete = lambda p, h=A: requests.delete(B + p, headers=h)

users = {u["username"]: u for u in get("/users?search=priya&page_size=100").json()["data"]}
priya = users["priya"]["id"]

print("== departments")
cse = check("create CSE", post("/departments", {"name": "Computer Science " + SFX, "code": "cs" + SFX}), 201)["data"]
expect("code upper-cased", cse["code"] == "CS" + SFX, cse["code"])
ece = check("create ECE", post("/departments", {"name": "Electronics " + SFX, "code": "EC" + SFX}), 201)["data"]
check("duplicate code", post("/departments", {"name": "X", "code": "CS" + SFX}), 409)
check("missing name", post("/departments", {"code": "ZZ" + SFX}), 400)
check("bad hod id", put(f"/departments/{cse['id']}", {"name": cse["name"], "code": cse["code"], "hod_id": 999999}), 400)
d = check("set HOD = priya", put(f"/departments/{cse['id']}", {"name": cse["name"], "code": cse["code"], "hod_id": priya}), 200)["data"]
expect("hod_name returned", d["hod_name"] == "Priya Raman", d["hod_name"])
lst = check("search departments", get(f"/departments?search=EC{SFX}"), 200)
expect("search finds 1", lst["meta"]["total"] == 1)

print("== academic years")
years = check("list years", get("/academic-years?all=true"), 200)["data"]
cur = next(y for y in years if y["is_current"])
check("end before start", post("/academic-years", {"name": "X" + SFX, "start_date": "2030-06-01", "end_date": "2030-01-01"}), 400)
ny = check("create current year", post("/academic-years", {"name": "T" + SFX, "start_date": "2040-06-01", "end_date": "2041-05-31", "is_current": True}), 201)["data"]
years = get("/academic-years?all=true").json()["data"]
expect("only one current", sum(y["is_current"] for y in years) == 1 and next(y for y in years if y["id"] == ny["id"])["is_current"])
check("restore original current", put(f"/academic-years/{cur['id']}", {**{k: cur[k] for k in ("name", "start_date", "end_date")}, "is_current": True}), 200)
check("delete test year", delete(f"/academic-years/{ny['id']}"), 200)
check("cannot delete current year", delete(f"/academic-years/{cur['id']}"), 409)

print("== subjects, exam types, levels, semesters")
s1 = check("create subject", post("/subjects", {"code": "cs3" + SFX + "1", "name": "Data Structures", "credits": 4, "subject_type": "theory"}), 201)["data"]
s2 = check("create lab", post("/subjects", {"code": "CS3" + SFX + "2", "name": "DS Lab", "credits": 2, "subject_type": "lab"}), 201)["data"]
check("negative credits", post("/subjects", {"code": "Q" + SFX, "name": "Q", "credits": -1}), 400)
check("bad subject_type", post("/subjects", {"code": "Q" + SFX, "name": "Q", "subject_type": "sports"}), 400)
check("exam types", get("/exam-types"), 200)
levels = check("year levels", get("/year-levels"), 200)["data"]
sems = check("semesters", get("/semesters"), 200)["data"]
check("year levels read-only", post("/year-levels", {"name": "V", "level_no": 5}), 404)
lvl2 = next(l for l in levels if l["level_no"] == 2)
sem3 = next(s for s in sems if s["sem_no"] == 3)

print("== curriculum")
r = check("add subjects to CSE sem 3", post("/curriculum", {"department_id": cse["id"], "semester_id": sem3["id"], "subject_ids": [s1["id"], s2["id"]]}), 201)
expect("added 2", r["data"]["added"] == 2)
r = check("re-add is idempotent", post("/curriculum", {"department_id": cse["id"], "semester_id": sem3["id"], "subject_ids": [s1["id"]]}), 201)
expect("added 0", r["data"]["added"] == 0)
cur_rows = check("list curriculum", get(f"/curriculum?department_id={cse['id']}"), 200)["data"]
expect("2 subjects", len(cur_rows) == 2)
check("subject in curriculum can't be deleted", delete(f"/subjects/{s1['id']}"), 409)
check("remove from curriculum", delete(f"/curriculum/{cur_rows[1]['id']}"), 200)

print("== staff")
kav = check("create staff (staff+hod)", post("/staff", {"name": "Kavya " + SFX, "employee_code": "emp" + SFX, "mobile": "9100000" + SFX,
                                                         "department_id": cse["id"], "designation": "Assistant Professor", "roles": ["staff", "hod"],
                                                         "password": "Staff@1234"}), 201)["data"]
expect("username from employee code", kav["username"] == "emp" + SFX, kav["username"])
check("bad staff role", post("/staff", {"name": "X", "employee_code": "Z" + SFX, "roles": ["admin"]}), 400)
check("duplicate employee code", post("/staff", {"name": "X", "employee_code": "EMP" + SFX}), 409)
p = check("give priya a staff profile + dept", put(f"/staff/{priya}", {"name": "Priya Raman", "employee_code": "EMP001", "username": "priya",
                                                                         "mobile": "9000000001", "department_id": cse["id"], "roles": ["staff", "placement_officer"]}), 200)["data"]
expect("priya roles kept", sorted(p["roles"]) == ["placement_officer", "staff"], p["roles"])

print("== classes")
c2a = check("create CSE II-A (incharge kavya)", post("/classes", {"department_id": cse["id"], "academic_year_id": cur["id"], "year_level_id": lvl2["id"], "section": "a", "class_incharge_id": kav["id"]}), 201)["data"]
expect("label", c2a["label"] == f"CS{SFX} II-A", c2a["label"])
check("duplicate class", post("/classes", {"department_id": cse["id"], "academic_year_id": cur["id"], "year_level_id": lvl2["id"], "section": "A"}), 409)
check("kavya can't be incharge twice", post("/classes", {"department_id": cse["id"], "academic_year_id": cur["id"], "year_level_id": lvl2["id"], "section": "C", "class_incharge_id": kav["id"]}), 409)
arun = get("/users?role=student&page_size=1").json()["data"][0]["id"]  # any student account
check("student can't be incharge", post("/classes", {"department_id": cse["id"], "academic_year_id": cur["id"], "year_level_id": lvl2["id"], "section": "C", "class_incharge_id": arun}), 400)
e1 = check("create ECE I-A", post("/classes", {"department_id": ece["id"], "academic_year_id": cur["id"], "year_level_id": levels[0]["id"], "section": "A"}), 201)["data"]
st = get(f"/staff?search=emp{SFX}").json()["data"][0]
expect("staff list shows incharge_of", st["incharge_of"] == c2a["label"], st["incharge_of"])

print("== students & parents")
fm = "9200000" + SFX
s_a = check("create student + 2 parents", post("/students", {
    "name": "Anu " + SFX, "register_no": "r" + SFX + "01", "department_id": cse["id"], "admission_year": 2025, "class_id": c2a["id"],
    "mobile": "9300000" + SFX, "gender": "female", "dob": "2007-04-01", "password": "Student@123",
    "parents": [{"name": "Father " + SFX, "mobile": fm, "relation": "father"}, {"name": "Mother " + SFX, "mobile": "9400000" + SFX, "relation": "mother"}]}), 201)["data"]
expect("username = register no", s_a["username"] == f"r{SFX}01", s_a["username"])
expect("batch defaulted", s_a["batch"] == "2025-2029", s_a["batch"])
expect("2 parents, father primary", len(s_a["parents"]) == 2 and s_a["parents"][0]["relation"] == "father" and s_a["parents"][0]["is_primary"])
s_b = check("sibling reuses father account", post("/students", {
    "name": "Balu " + SFX, "register_no": "R" + SFX + "02", "department_id": cse["id"], "admission_year": 2025, "class_id": c2a["id"],
    "parents": [{"name": "Father " + SFX, "mobile": fm, "relation": "father"}]}), 201)["data"]
expect("same parent id", s_b["parents"][0]["id"] == s_a["parents"][0]["id"])
check("class of another department", post("/students", {"name": "X", "register_no": "X" + SFX, "department_id": cse["id"], "admission_year": 2025, "class_id": e1["id"]}), 400)
check("duplicate register no", post("/students", {"name": "X", "register_no": "R" + SFX + "01", "department_id": cse["id"], "admission_year": 2025}), 409)
check("bad parent mobile", post("/students", {"name": "X", "register_no": "Y" + SFX, "department_id": cse["id"], "admission_year": 2025, "parents": [{"name": "P", "mobile": "123"}]}), 400)
s_c = check("ECE student", post("/students", {"name": "Chitra " + SFX, "register_no": "R" + SFX + "03", "department_id": ece["id"], "admission_year": 2026, "class_id": e1["id"]}), 201)["data"]
lst = check("filter by class", get(f"/students?class_id={c2a['id']}"), 200)
expect("2 in CSE II-A", lst["meta"]["total"] == 2)
cls = get(f"/classes/{c2a['id']}").json()["data"]
expect("class student_count", cls["student_count"] == 2)
g = check("add guardian to ECE student", post(f"/students/{s_c['id']}/parents", {"name": "Uncle " + SFX, "mobile": "9500000" + SFX, "relation": "guardian"}), 200)["data"]
check("link same parent twice", post(f"/students/{s_c['id']}/parents", {"id": g["parents"][0]["id"], "relation": "guardian"}), 409)
check("unlink guardian", delete(f"/students/{s_c['id']}/parents/{g['parents'][0]['id']}"), 200)
upd = check("update student (move section)", put(f"/students/{s_b['id']}", {"name": "Balu " + SFX, "register_no": "R" + SFX + "02", "department_id": cse["id"], "admission_year": 2025, "class_id": c2a["id"], "blood_group": "O+"}), 200)["data"]
expect("blood group saved", upd["blood_group"] == "O+")
check("parent directory", get(f"/parents?search={fm}"), 200)

print("== data scoping")
father = otp_login(s_a["parents"][0]["username"])
fl = check("father lists students", get("/students", father), 200)
expect("father sees exactly his 2 children", sorted(s["id"] for s in fl["data"]) == sorted([s_a["id"], s_b["id"]]), fl["meta"]["total"])
check("father can't open another student", get(f"/students/{s_c['id']}", father), 404)
check("father can't create students", post("/students", {"name": "X"}, father), 403)
anu = login(s_a["username"], "Student@123")
al = check("student lists students", get("/students", anu), 200)
expect("student sees only herself", [s["id"] for s in al["data"]] == [s_a["id"]])
pr = login("priya", "Staff@1234")
# Search by this run's register number: the test database grows past one page over time.
pa = check("placement officer lists students", get(f"/students?search=R{SFX}&page_size=100", pr), 200)
expect("placement officer sees every department", s_c["id"] in [s["id"] for s in pa["data"]])
kv = login(kav["username"], "Staff@1234")
pl = check("HOD/staff (CSE) lists students", get("/students?page_size=100", kv), 200)
expect("staff sees only CSE", all(s["department_id"] == cse["id"] for s in pl["data"]) and s_c["id"] not in [s["id"] for s in pl["data"]])
check("CSE staff can't open ECE student", get(f"/students/{s_c['id']}", kv), 404)

print("== bulk upload")
t = get("/bulk-upload/students/template")
expect("template is xlsx", t.status_code == 200 and t.content[:2] == b"PK", t.status_code)
headers = [c.value for c in next(openpyxl.load_workbook(io.BytesIO(t.content)).active.iter_rows(max_row=1))]
expect("template headers", headers[:3] == ["register_no*", "name*", "department_code*"], headers[:3])

wb = openpyxl.Workbook()
ws = wb.active
ws.append(["name*", "register_no*", "department_code*", "admission_year*", "year*", "section", "mobile", "dob", "parent_name", "parent_mobile", "parent_relation"])
ws.append(["Bulk One", "B" + SFX + "01", "cs" + SFX, 2025, 2, "A", 9600000000 + int(SFX), "2007-01-15", "BP One", fm, "father"])  # sibling via father mobile
ws.append(["Bulk Two", "B" + SFX + "02", "CS" + SFX, 2025, 2, "B", None, "15-02-2007", "BP Two", "9700000" + SFX, "mother"])  # section B missing
ws.append(["Bad Dept", "B" + SFX + "03", "NOPE", 2025, 2, "A"])
ws.append(["Bad Year", "B" + SFX + "04", "CS" + SFX, 2025, 7, "A"])
ws.append(["Dup Reg", "R" + SFX + "01", "CS" + SFX, 2025, 2, "A"])
ws.append([None] * 5)  # blank row ignored
ws.append(["Bad Parent", "B" + SFX + "05", "CS" + SFX, 2025, 2, "A", None, None, "No Mobile", "12", "father"])
buf = io.BytesIO()
wb.save(buf)
xlsx = buf.getvalue()
up = lambda **form: requests.post(B + "/bulk-upload/students", headers=A, files={"file": ("students.xlsx", xlsx)}, data=form)

before = get("/students?page_size=1").json()["meta"]["total"]
j = check("dry run", up(dry_run="true"), 201)["data"]
expect("dry run: 1 ok, 5 errors", j["success_rows"] == 1 and j["failed_rows"] == 5, f"{j['success_rows']}/{j['failed_rows']} {[e['message'] for e in j['errors']]}")
expect("dry run saved nothing", get("/students?page_size=1").json()["meta"]["total"] == before)
j = check("real run, create missing classes", up(create_missing_classes="true", default_password="Welcome@123"), 201)["data"]
expect("2 ok, 4 errors", j["success_rows"] == 2 and j["failed_rows"] == 4, f"{j['success_rows']}/{j['failed_rows']}")
msgs = {e["row"]: e["message"] for e in j["errors"]}
expect("row errors are specific", "department_code" in msgs.get(4, "") and "year must" in msgs.get(5, "") and "register number" in msgs.get(6, "").lower(), msgs)
expect("blank row skipped; bad parent reported on row 8", "Parent 1" in msgs.get(8, ""), msgs.get(8))
b1 = get(f"/students?search=B{SFX}01").json()["data"][0]
b1d = get(f"/students/{b1['id']}").json()["data"]
expect("bulk sibling linked to existing father", b1d["parents"][0]["id"] == s_a["parents"][0]["id"])
expect("mobile number kept intact", b1d["mobile"] == str(9600000000 + int(SFX)), b1d["mobile"])
b2 = get(f"/students?search=B{SFX}02").json()["data"][0]
expect("class II-B auto-created", b2["class_label"] == f"CS{SFX} II-B", b2["class_label"])
expect("DD-MM-YYYY dob parsed", get(f"/students/{b2['id']}").json()["data"]["dob"] == "2007-02-15")
check("default password works", requests.post(f"{B}/auth/login", json={"username": b1["username"], "password": "Welcome@123"}), 200)
check("upload history", get("/bulk-upload/jobs"), 200)
rep = get(f"/bulk-upload/jobs/{j['id']}/errors.xlsx")
expect("error report xlsx", rep.status_code == 200 and rep.content[:2] == b"PK")
bad = requests.post(B + "/bulk-upload/students", headers=A, files={"file": ("x.xlsx", b"not excel")})
check("garbage file rejected", bad, 400)
check("staff can't bulk upload", requests.post(B + "/bulk-upload/students", headers=kv, files={"file": ("s.xlsx", xlsx)}), 403)

print("== delete guards")
check("class with students can't be deleted", delete(f"/classes/{c2a['id']}"), 409)
check("department in use can't be deleted", delete(f"/departments/{cse['id']}"), 409)
check("incharge can't be deactivated", patch(f"/staff/{kav['id']}/status", {"is_active": False}), 409)

print(f"\nPASSED={PASS} FAILED={FAIL}")
sys.exit(1 if FAIL else 0)
